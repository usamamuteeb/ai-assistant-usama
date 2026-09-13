from __future__ import annotations

from pathlib import Path

import pytest

from plugins.knowledge_base.plugin import (
    AskDocumentTool,
    KnowledgeBaseSearchTool,
    KnowledgeBaseTool,
    ListKnowledgeDocumentsTool,
    RemoveKnowledgeDocumentTool,
    SearchKnowledgeBaseWithCitationsTool,
    TagDocumentTool,
)
import plugins.knowledge_base.plugin as kb


class State:
    def __init__(self):
        self.values = {}

    def get_state(self, key, default=None):
        return self.values.get(key, default)

    def set_state(self, key, value):
        self.values[key] = value


class Vector:
    def __init__(self):
        self.items = []

    def add_memory(self, text, metadata=None):
        self.items.append({"text": text, "metadata": metadata or {}})


def make_ingest(folder: Path):
    tool = object.__new__(KnowledgeBaseTool)
    tool.knowledge_base_dir = folder
    tool.store = State()
    tool.vector_memory = Vector()
    tool.chunk_size = 20
    tool.chunk_overlap = 0
    return tool


def test_pdf_chunks_keep_page_numbers_and_content_hash(monkeypatch, tmp_path):
    pdf = tmp_path / "report.pdf"
    pdf.write_bytes(b"fixture")
    tool = make_ingest(tmp_path)
    monkeypatch.setattr(
        kb,
        "_extract_document",
        lambda path: {
            "text": "page one\n\npage two",
            "segments": [
                {"text": "page one", "metadata": {"page_number": 1}},
                {"text": "page two", "metadata": {"page_number": 2}},
            ],
            "page_count": 2,
        },
    )

    result = tool.run()

    assert result["ingested_files"] == ["report.pdf"]
    assert [item["metadata"]["page_number"] for item in tool.vector_memory.items] == [1, 2]
    record = tool.store.get_state("kb_file_hashes")["report.pdf"]
    assert record["page_count"] == 2
    assert len(record["content_hash"]) == 64


def test_pdf_extractor_reads_each_page_independently(monkeypatch, tmp_path):
    class FakePage:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    class FakeReader:
        def __init__(self, _path):
            self.pages = [FakePage("first page"), FakePage("second page")]

    import pypdf
    monkeypatch.setattr(pypdf, "PdfReader", FakeReader)
    path = tmp_path / "pages.pdf"
    path.write_bytes(b"fixture")

    document = kb._extract_document(path)

    assert document["page_count"] == 2
    assert [segment["metadata"]["page_number"] for segment in document["segments"]] == [1, 2]
    assert document["text"] == "first page\n\nsecond page"


def test_xlsx_extractor_keeps_sheet_metadata(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Name", "Value"])
    sheet.append(["Latency", 42])
    path = tmp_path / "metrics.xlsx"
    workbook.save(path)
    workbook.close()

    document = kb._extract_document(path)

    assert document["sheet_count"] == 1
    assert document["segments"][0]["metadata"]["sheet_name"] == "Metrics"
    assert "Latency\t42" in document["segments"][0]["text"]


def test_same_extracted_content_under_new_name_is_reported_as_duplicate(tmp_path):
    original = tmp_path / "original.txt"
    original.write_text("same content", encoding="utf-8")
    tool = make_ingest(tmp_path)
    assert tool.run()["ingested_files"] == ["original.txt"]

    duplicate = tmp_path / "renamed.txt"
    duplicate.write_text("same content", encoding="utf-8")
    result = tool.run()

    assert result["duplicates"] == [{"file": "renamed.txt", "duplicate_of": "original.txt"}]
    assert result["total_chunks_added"] == 0
    assert "renamed.txt" not in tool.store.get_state("kb_file_hashes")


def test_docx_and_pptx_extraction_is_supported(tmp_path):
    docx = pytest.importorskip("docx")
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    docx_path = tmp_path / "notes.docx"
    document = docx.Document()
    document.add_paragraph("DOCX searchable text")
    document.save(docx_path)

    pptx_path = tmp_path / "slides.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "PPTX searchable text"
    presentation.save(pptx_path)

    tool = make_ingest(tmp_path)
    result = tool.run()

    assert set(result["ingested_files"]) == {"notes.docx", "slides.pptx"}
    assert {item["text"] for item in tool.vector_memory.items} == {"DOCX searchable text", "PPTX searchable text"}


def test_tag_list_and_reingest_preserve_metadata(tmp_path):
    path = tmp_path / "tagged.md"
    path.write_text("initial text", encoding="utf-8")
    tool = make_ingest(tmp_path)
    tool.run()
    tagger = TagDocumentTool(tool.store, tool.vector_memory)

    assert tagger.run("tagged.md", tags=["work", "important"], category="engineering")["category"] == "engineering"
    assert tagger.run("tagged.md", category="engineering")["tags"] == ["work", "important"]
    path.write_text("updated text", encoding="utf-8")
    tool.run()

    listed = ListKnowledgeDocumentsTool(tool.store).run()["documents"][0]
    assert listed["tags"] == ["work", "important"]
    assert listed["category"] == "engineering"


def test_remove_requires_confirmation_only_when_deleting_file(tmp_path):
    path = tmp_path / "remove.md"
    path.write_text("remove me", encoding="utf-8")
    tool = make_ingest(tmp_path)
    tool.run()
    confirm_calls = []
    remover = RemoveKnowledgeDocumentTool(tool.store, tool.vector_memory, tmp_path, lambda prompt: confirm_calls.append(prompt) or False)

    assert remover.run("remove.md") == {"removed": "remove.md", "file_deleted": False}
    path.write_text("remove me", encoding="utf-8")
    tool.run()
    assert remover.run("remove.md", delete_file=True) == {"error": "Action not performed: confirmation denied or not provided."}
    assert confirm_calls
    assert path.exists()


class QueryCollection:
    def __init__(self, entries):
        self.entries = entries

    def query(self, query_texts, n_results, where):
        name = next(iter(where.values()))
        matches = [entry for entry in self.entries if entry["metadata"].get("source_file") == name]
        return {
            "documents": [[entry["text"] for entry in matches[:n_results]]],
            "metadatas": [[entry["metadata"] for entry in matches[:n_results]]],
            "distances": [[0.1 for _ in matches[:n_results]]],
        }


class QueryVector:
    def __init__(self, entries):
        self.collection = QueryCollection(entries)

    def search(self, query, k=5):
        return []


def test_filtered_search_and_citations_are_scoped():
    entries = [
        {"text": "engineering result", "metadata": {"source_file": "eng.md", "category": "engineering", "ingested_at": "2026-09-10T00:00:00+00:00", "page_number": 2}},
        {"text": "finance result", "metadata": {"source_file": "finance.md", "category": "finance", "ingested_at": "2026-09-10T00:00:00+00:00"}},
    ]
    vector = QueryVector(entries)
    state = State()
    state.set_state("kb_file_hashes", {
        "eng.md": {"content_hash": "a", "category": "engineering", "ingested_at": "2026-09-10T00:00:00+00:00"},
        "finance.md": {"content_hash": "b", "category": "finance", "ingested_at": "2026-09-10T00:00:00+00:00"},
    })

    result = SearchKnowledgeBaseWithCitationsTool(vector, state).run("result", category="engineering")
    assert result["results"][0]["citation"] == "eng.md, page 2"
    assert result["results"][0]["metadata"]["source_file"] == "eng.md"


def test_ask_document_returns_only_requested_file():
    entries = [
        {"text": "one", "metadata": {"source_file": "one.md", "page_number": 1}},
        {"text": "two", "metadata": {"source_file": "two.md", "page_number": 1}},
    ]
    result = AskDocumentTool(QueryVector(entries)).run("one.md", "question")
    assert [item["text"] for item in result["results"]] == ["one"]
    assert result["results"][0]["page_number"] == 1
