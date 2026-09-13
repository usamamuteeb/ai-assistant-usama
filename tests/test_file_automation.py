"""Focused tests for the file and document automation plugin."""
from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

from fpdf import FPDF
from docx import Document

from plugins.file_automation.plugin import (
    CopyFileTool,
    CreateFolderTool,
    CreateMarkdownReportTool,
    CreatePdfReportTool,
    CreateWordDocumentTool,
    ExtractPdfTablesTool,
    MoveFileTool,
    RenameFileTool,
    _context,
)


class ConfirmationBridge:
    def __init__(self, result: bool = True):
        self.calls: list[str] = []
        self.result = result

    def __call__(self, message: str) -> bool:
        self.calls.append(message)
        return self.result

    def manual_confirm(self, message: str) -> bool:
        self.calls.append(f"manual: {message}")
        return self.result


def _context_for(tmp_path: Path, confirm: ConfirmationBridge | None = None):
    config = tmp_path / "config.yaml"
    config.write_text(
        """filesystem_tool:
  workspace_root: workspace
dev_tools:
  protected_path_prefixes:
    - 'C:\\Windows'
file_automation:
  force_manual_confirmation: true
  backup_folder: file_backups
""",
        encoding="utf-8",
    )
    return _context(tmp_path, confirm or ConfirmationBridge())


def test_reports_create_expected_markdown_and_word_content(tmp_path):
    context = _context_for(tmp_path)
    body = "## Summary\n\nThe deployment is ready."

    markdown = CreateMarkdownReportTool(context).run(str(tmp_path / "report"), "Deployment", body)
    word = CreateWordDocumentTool(context).run(str(tmp_path / "report"), "Deployment", body)

    assert markdown["status"] == "created"
    assert "# Deployment" in Path(markdown["path"]).read_text(encoding="utf-8")
    document = Document(word["path"])
    assert "Deployment" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "The deployment is ready." in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_pdf_report_and_table_extraction(tmp_path):
    context = _context_for(tmp_path)
    report = CreatePdfReportTool(context).run(
        str(tmp_path / "report.pdf"), "Deployment", "The deployment is ready."
    )
    assert report["status"] == "created"
    assert Path(report["path"]).stat().st_size > 0

    table_path = tmp_path / "table.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for row in (("Name", "Status"), ("API", "Ready")):
        for cell in row:
            pdf.cell(70, 10, cell, border=1)
        pdf.ln()
    pdf.output(str(table_path))

    result = ExtractPdfTablesTool(context).run(str(table_path))
    assert result["tables"]
    assert result["tables"][0]["rows"][:2] == [["Name", "Status"], ["API", "Ready"]]


def test_overwrite_is_forced_manual_and_backed_up(tmp_path):
    confirm = ConfirmationBridge()
    context = _context_for(tmp_path, confirm)
    source = tmp_path / "source.txt"
    destination = tmp_path / "destination.txt"
    source.write_text("new content", encoding="utf-8")
    destination.write_text("old content", encoding="utf-8")

    result = MoveFileTool(context).run(str(source), str(destination))

    assert result["status"] == "moved"
    assert destination.read_text(encoding="utf-8") == "new content"
    backup = Path(result["backup"])
    assert backup.exists()
    assert backup.read_text(encoding="utf-8") == "old content"
    assert any(message.startswith("manual:") for message in confirm.calls)


def test_copy_new_destination_is_additive_but_overwrite_is_manual(tmp_path):
    confirm = ConfirmationBridge()
    context = _context_for(tmp_path, confirm)
    source = tmp_path / "source.txt"
    destination = tmp_path / "copy.txt"
    source.write_text("content", encoding="utf-8")

    created = CopyFileTool(context).run(str(source), str(destination))
    assert created["status"] == "copied"
    assert confirm.calls == []

    overwritten = CopyFileTool(context).run(str(source), str(destination))
    assert overwritten["status"] == "copied"
    assert overwritten["backup"]
    assert any(message.startswith("manual:") for message in confirm.calls)


def test_rename_dry_run_does_not_change_file(tmp_path):
    context = _context_for(tmp_path)
    source = tmp_path / "before.txt"
    source.write_text("keep", encoding="utf-8")

    result = RenameFileTool(context).run(str(source), "after.txt", dry_run=True)

    assert result["dry_run"] is True
    assert result["would_overwrite"] is False
    assert source.exists()
    assert not (tmp_path / "after.txt").exists()


def test_protected_path_is_refused_before_confirmation(tmp_path):
    confirm = MagicMock(return_value=True)
    context = _context_for(tmp_path, confirm)

    result = CreateFolderTool(context).run(r"C:\Windows\system32\assistant-test")

    assert result == {"error": "Refused: this path is in the protected-paths blocklist in config.yaml."}
    confirm.assert_not_called()
